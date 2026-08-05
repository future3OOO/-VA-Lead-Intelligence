from __future__ import annotations

import csv
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

LEAD_FIELDS = [
    "lead_id",
    "company_id",
    "primary_contact_id",
    "company_name",
    "primary_domain",
    "primary_contact_name",
    "primary_contact_title",
    "primary_contact_email",
    "primary_contact_phone",
    "primary_contact_linkedin",
    "company_email",
    "company_phone",
    "company_form",
    "best_email",
    "best_phone",
    "best_form",
    "job_title",
    "location",
    "workplace_type",
    "category",
    "source",
    "source_url",
    "intent_label",
    "published_at",
    "qualification_score",
    "rank",
    "explanation",
]

PRIMARY_CONTACT_FIELDS = [
    "lead_id",
    "company_id",
    "primary_contact_id",
    "company_name",
    "primary_domain",
    "job_title",
    "primary_contact_status",
    "primary_contact_name",
    "primary_contact_title",
    "primary_contact_email",
    "primary_contact_phone",
    "primary_contact_linkedin",
    "source",
    "source_url",
]

CONTACT_FIELDS = [
    "contact_id",
    "company_id",
    "company_name",
    "primary_domain",
    "contact_name",
    "contact_title",
    "contact_emails",
    "contact_phones",
    "contact_linkedin_urls",
]

COMPANY_FIELDS = [
    "company_id",
    "company_name",
    "primary_domain",
    "target_fit",
    "source_hit_count",
    "best_email",
    "best_phone",
    "best_form",
    "named_contact_name",
    "named_contact_title",
    "named_contact_email",
    "named_contact_phone",
    "named_contact_linkedin",
]


def _write_csv(path: Path, fields: list[str], row: dict[str, str] | None) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        if row is not None:
            writer.writerow(row)


def _distinct_row(fields: list[str], prefix: str) -> dict[str, str]:
    row = {field: f"{prefix}:{field}" for field in fields}
    for index, field in enumerate(fields):
        if "phone" in field:
            row[field] = f"+64 21 555 {index:04d}"
        elif "email" in field:
            row[field] = f"{prefix}.{field}@example.org"
        elif "linkedin" in field:
            row[field] = f"https://www.linkedin.com/in/{prefix}-{field}"
    return row


def test_workbook_cli_preserves_all_four_relational_export_tables(tmp_path: Path) -> None:
    leads_path = tmp_path / "leads.csv"
    primary_contacts_path = tmp_path / "primary_contacts.csv"
    contacts_path = tmp_path / "contacts.csv"
    companies_path = tmp_path / "companies.csv"
    workbook_path = tmp_path / "leads.xlsx"

    lead_row = _distinct_row(LEAD_FIELDS, "lead")
    primary_contact_row = _distinct_row(PRIMARY_CONTACT_FIELDS, "primary")
    contact_row = _distinct_row(CONTACT_FIELDS, "contact")
    company_row = _distinct_row(COMPANY_FIELDS, "company")
    lead_row.update(
        lead_id="lead-1",
        company_id="company-1",
        primary_contact_id="contact-1",
    )
    primary_contact_row.update(
        lead_id="lead-1",
        company_id="company-1",
        primary_contact_id="contact-1",
        primary_contact_status="available",
    )
    contact_row.update(contact_id="contact-1", company_id="company-1")
    company_row["company_id"] = "company-1"
    _write_csv(leads_path, LEAD_FIELDS, lead_row)
    _write_csv(primary_contacts_path, PRIMARY_CONTACT_FIELDS, primary_contact_row)
    _write_csv(contacts_path, CONTACT_FIELDS, contact_row)
    _write_csv(companies_path, COMPANY_FIELDS, company_row)

    result = subprocess.run(
        [
            sys.executable,
            "scripts/build_leads_workbook.py",
            "--leads",
            str(leads_path),
            "--primary-contacts",
            str(primary_contacts_path),
            "--contacts",
            str(contacts_path),
            "--companies",
            str(companies_path),
            "--output",
            str(workbook_path),
        ],
        cwd=Path(__file__).resolve().parents[2],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    assert workbook.sheetnames == ["Leads", "Primary Contacts", "Contacts", "Companies"]

    expected = [
        ("Leads", "LeadsTable"),
        ("Primary Contacts", "PrimaryContactsTable"),
        ("Contacts", "ContactsTable"),
        ("Companies", "CompaniesTable"),
    ]
    for sheet_name, table_name in expected:
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        assert len(headers) == len(set(headers))
        assert sheet.max_row == 2
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is None
        assert sheet.row_dimensions[1].height == 18
        assert all(not cell.alignment.wrap_text for cell in sheet[1])
        assert all(not cell.alignment.wrap_text for cell in sheet[2])
        table = sheet.tables[table_name]
        assert table.ref == sheet.dimensions
        assert table.autoFilter.ref == sheet.dimensions
        assert sheet.tables[table_name].tableStyleInfo is None

    assert [cell.value for cell in workbook["Leads"][1]][:9] == [
        "Company",
        "Domain",
        "Lead title",
        "Contact role",
        "Contact name",
        "Contact title",
        "Email(s)",
        "Phone(s)",
        "LinkedIn profile(s)",
    ]
    assert workbook["Leads"]["E2"].value == contact_row["contact_name"]
    assert workbook["Primary Contacts"]["E2"].value == primary_contact_row["primary_contact_name"]
    assert workbook["Contacts"]["C2"].value == contact_row["contact_name"]
    assert workbook["Companies"]["A2"].value == company_row["company_name"]
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            if cell.value in {"Lead ID", "Company ID", "Contact ID", "Primary contact ID"}:
                assert sheet.column_dimensions[cell.column_letter].hidden
    primary_sheet = workbook["Primary Contacts"]
    for header in (
        "Selected primary name",
        "Selected primary title",
        "Selected primary email",
        "Selected primary phone",
        "Selected primary LinkedIn",
    ):
        cell = next(cell for cell in primary_sheet[1] if cell.value == header)
        assert not primary_sheet.column_dimensions[cell.column_letter].hidden
    workbook_values = {
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(min_row=2)
        for cell in row
        if cell.value is not None
    }
    for source_row in (lead_row, primary_contact_row, contact_row, company_row):
        assert {str(value) for value in source_row.values() if value} <= workbook_values

    with zipfile.ZipFile(workbook_path) as package:
        workbook_xml = package.read("xl/workbook.xml")
        assert b"_FilterDatabase" not in workbook_xml
        for index in range(1, 5):
            worksheet_xml = package.read(f"xl/worksheets/sheet{index}.xml")
            table_xml = package.read(f"xl/tables/table{index}.xml")
            assert b"<autoFilter" not in worksheet_xml
            assert b"<autoFilter" in table_xml


@pytest.mark.parametrize(
    "invalid_input",
    ["leads", "primary_contacts", "contacts", "companies"],
)
def test_workbook_cli_rejects_an_unexpected_csv_contract(
    tmp_path: Path, invalid_input: str
) -> None:
    leads_path = tmp_path / "leads.csv"
    primary_contacts_path = tmp_path / "primary_contacts.csv"
    contacts_path = tmp_path / "contacts.csv"
    companies_path = tmp_path / "companies.csv"
    workbook_path = tmp_path / "leads.xlsx"
    _write_csv(leads_path, LEAD_FIELDS, {})
    _write_csv(primary_contacts_path, PRIMARY_CONTACT_FIELDS, {})
    _write_csv(contacts_path, CONTACT_FIELDS, {})
    _write_csv(companies_path, COMPANY_FIELDS, {})
    paths = {
        "leads": leads_path,
        "primary_contacts": primary_contacts_path,
        "contacts": contacts_path,
        "companies": companies_path,
    }
    _write_csv(paths[invalid_input], ["wrong_column"], {"wrong_column": "value"})

    result = subprocess.run(
        [
            sys.executable,
            "scripts/build_leads_workbook.py",
            "--leads",
            str(leads_path),
            "--primary-contacts",
            str(primary_contacts_path),
            "--contacts",
            str(contacts_path),
            "--companies",
            str(companies_path),
            "--output",
            str(workbook_path),
        ],
        cwd=Path(__file__).resolve().parents[2],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 2
    assert "unexpected columns" in result.stderr
    assert not workbook_path.exists()


def test_workbook_cli_rejects_mixed_exports_with_broken_relations(tmp_path: Path) -> None:
    leads_path = tmp_path / "leads.csv"
    primary_contacts_path = tmp_path / "primary_contacts.csv"
    contacts_path = tmp_path / "contacts.csv"
    companies_path = tmp_path / "companies.csv"
    workbook_path = tmp_path / "leads.xlsx"

    _write_csv(
        leads_path,
        LEAD_FIELDS,
        {"lead_id": "lead-1", "company_id": "company-1", "primary_contact_id": "contact-1"},
    )
    _write_csv(
        primary_contacts_path,
        PRIMARY_CONTACT_FIELDS,
        {"lead_id": "lead-1", "company_id": "company-1", "primary_contact_id": "contact-2"},
    )
    _write_csv(
        contacts_path,
        CONTACT_FIELDS,
        {"contact_id": "contact-1", "company_id": "company-1"},
    )
    _write_csv(companies_path, COMPANY_FIELDS, {"company_id": "company-1"})

    result = subprocess.run(
        [
            sys.executable,
            "scripts/build_leads_workbook.py",
            "--leads",
            str(leads_path),
            "--primary-contacts",
            str(primary_contacts_path),
            "--contacts",
            str(contacts_path),
            "--companies",
            str(companies_path),
            "--output",
            str(workbook_path),
        ],
        cwd=Path(__file__).resolve().parents[2],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 2
    assert "lead and primary contact references differ" in result.stderr
    assert not workbook_path.exists()


def test_workbook_cli_rejects_an_inconsistent_primary_status(tmp_path: Path) -> None:
    leads_path = tmp_path / "leads.csv"
    primary_contacts_path = tmp_path / "primary_contacts.csv"
    contacts_path = tmp_path / "contacts.csv"
    companies_path = tmp_path / "companies.csv"
    workbook_path = tmp_path / "leads.xlsx"
    _write_csv(leads_path, LEAD_FIELDS, {"lead_id": "lead-1", "company_id": "company-1"})
    _write_csv(
        primary_contacts_path,
        PRIMARY_CONTACT_FIELDS,
        {
            "lead_id": "lead-1",
            "company_id": "company-1",
            "primary_contact_status": "available",
        },
    )
    _write_csv(
        contacts_path,
        CONTACT_FIELDS,
        {"contact_id": "contact-other", "company_id": "company-1", "contact_name": "Other"},
    )
    _write_csv(companies_path, COMPANY_FIELDS, {"company_id": "company-1"})

    result = subprocess.run(
        [
            sys.executable,
            "scripts/build_leads_workbook.py",
            "--leads",
            str(leads_path),
            "--primary-contacts",
            str(primary_contacts_path),
            "--contacts",
            str(contacts_path),
            "--companies",
            str(companies_path),
            "--output",
            str(workbook_path),
        ],
        cwd=Path(__file__).resolve().parents[2],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 2
    assert "available without a contact ID and name" in result.stderr
    assert not workbook_path.exists()


def test_workbook_cli_accepts_a_preserved_unresolved_lead(tmp_path: Path) -> None:
    leads_path = tmp_path / "leads.csv"
    primary_contacts_path = tmp_path / "primary_contacts.csv"
    contacts_path = tmp_path / "contacts.csv"
    companies_path = tmp_path / "companies.csv"
    workbook_path = tmp_path / "leads.xlsx"
    _write_csv(leads_path, LEAD_FIELDS, {"lead_id": "lead-1", "company_name": "Unresolved"})
    _write_csv(
        primary_contacts_path,
        PRIMARY_CONTACT_FIELDS,
        {
            "lead_id": "lead-1",
            "company_name": "Unresolved",
            "primary_contact_status": "unavailable",
        },
    )
    _write_csv(contacts_path, CONTACT_FIELDS, None)
    _write_csv(companies_path, COMPANY_FIELDS, None)

    result = subprocess.run(
        [
            sys.executable,
            "scripts/build_leads_workbook.py",
            "--leads",
            str(leads_path),
            "--primary-contacts",
            str(primary_contacts_path),
            "--contacts",
            str(contacts_path),
            "--companies",
            str(companies_path),
            "--output",
            str(workbook_path),
        ],
        cwd=Path(__file__).resolve().parents[2],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    workbook = load_workbook(workbook_path, read_only=True)
    lead_headers = [cell.value for cell in workbook["Leads"][1]]
    primary_headers = [cell.value for cell in workbook["Primary Contacts"][1]]
    assert workbook["Leads"][2][lead_headers.index("Lead ID")].value == "lead-1"
    assert (
        workbook["Primary Contacts"][2][primary_headers.index("Contact status")].value
        == "unavailable"
    )
    assert workbook["Leads"].max_row == 2


def test_workbook_leads_exposes_every_company_contact_without_visible_mapping_ids(
    tmp_path: Path,
) -> None:
    leads_path = tmp_path / "leads.csv"
    primary_contacts_path = tmp_path / "primary_contacts.csv"
    contacts_path = tmp_path / "contacts.csv"
    companies_path = tmp_path / "companies.csv"
    workbook_path = tmp_path / "leads.xlsx"
    _write_csv(
        leads_path,
        LEAD_FIELDS,
        {
            "lead_id": "lead-1",
            "company_id": "company-1",
            "primary_contact_id": "contact-1",
            "company_name": "Alpha Realty",
            "primary_domain": "alpha.example.org",
            "primary_contact_name": "Alice Morgan",
            "job_title": "Property management support",
            "location": "'-43.5321, 172.6362, New Zealand",
        },
    )
    _write_csv(
        primary_contacts_path,
        PRIMARY_CONTACT_FIELDS,
        {
            "lead_id": "lead-1",
            "company_id": "company-1",
            "primary_contact_id": "contact-1",
            "company_name": "Alpha Realty",
            "primary_contact_status": "available",
            "primary_contact_name": "Alice Morgan",
        },
    )
    with contacts_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=CONTACT_FIELDS)
        writer.writeheader()
        writer.writerows(
            [
                {
                    "contact_id": "contact-1",
                    "company_id": "company-1",
                    "company_name": "Alpha Realty",
                    "contact_name": "Alice Morgan",
                    "contact_emails": "alice@alpha.example.org",
                },
                {
                    "contact_id": "contact-2",
                    "company_id": "company-1",
                    "company_name": "Alpha Realty",
                    "contact_name": "Bob Taylor",
                    "contact_phones": "+64 21 555 0102",
                    "contact_linkedin_urls": "https://www.linkedin.com/in/bob-taylor",
                },
            ]
        )
    _write_csv(
        companies_path,
        COMPANY_FIELDS,
        {"company_id": "company-1", "company_name": "Alpha Realty"},
    )

    result = subprocess.run(
        [
            sys.executable,
            "scripts/build_leads_workbook.py",
            "--leads",
            str(leads_path),
            "--primary-contacts",
            str(primary_contacts_path),
            "--contacts",
            str(contacts_path),
            "--companies",
            str(companies_path),
            "--output",
            str(workbook_path),
        ],
        cwd=Path(__file__).resolve().parents[2],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    sheet = workbook["Leads"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:9] == [
        "Company",
        "Domain",
        "Lead title",
        "Contact role",
        "Contact name",
        "Contact title",
        "Email(s)",
        "Phone(s)",
        "LinkedIn profile(s)",
    ]
    assert sheet.max_row == 3
    assert [sheet.cell(row=row, column=5).value for row in (2, 3)] == [
        "Alice Morgan",
        "Bob Taylor",
    ]
    assert [sheet.cell(row=row, column=4).value for row in (2, 3)] == [
        "Primary",
        "Additional",
    ]
    linkedin_cell = sheet.cell(row=3, column=9)
    assert linkedin_cell.value == "https://www.linkedin.com/in/bob-taylor"
    assert linkedin_cell.hyperlink is not None
    assert linkedin_cell.hyperlink.target == linkedin_cell.value
    location_column = headers.index("Location") + 1
    assert sheet.cell(row=2, column=location_column).value == "New Zealand"
    for field in ("Lead ID", "Company ID", "Contact ID", "Primary contact ID", "Coordinates"):
        column = headers.index(field) + 1
        letter = sheet.cell(row=1, column=column).column_letter
        assert sheet.column_dimensions[letter].hidden
